from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUTPUT = "python_learning_progress.pptx"


def add_text(slide, text, left, top, width, height, size, color, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return box


def add_card(slide, left, top, width, height, number, title, body, accent):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.color.rgb = RGBColor(226, 231, 238)
    card.line.width = Pt(1)

    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(left + 0.22), Inches(top + 0.24), Inches(0.42), Inches(0.42)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(*accent)
    badge.line.fill.background()
    badge.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_text(slide, number, left + 0.22, top + 0.285, 0.42, 0.22, 11,
             (255, 255, 255), True, align=PP_ALIGN.CENTER)
    add_text(slide, title, left + 0.78, top + 0.22, width - 1.0, 0.30, 14,
             (25, 39, 58), True)
    add_text(slide, body, left + 0.78, top + 0.62, width - 1.0, height - 0.78, 10.5,
             (83, 96, 112))


presentation = Presentation()
presentation.slide_width = Inches(13.333)
presentation.slide_height = Inches(7.5)
slide = presentation.slides.add_slide(presentation.slide_layouts[6])

background = slide.background.fill
background.solid()
background.fore_color.rgb = RGBColor(246, 248, 251)

header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.42))
header.fill.solid()
header.fill.fore_color.rgb = RGBColor(24, 49, 77)
header.line.fill.background()

add_text(slide, "Python Learning Progress", 0.62, 0.35, 8.7, 0.5, 27,
         (255, 255, 255), True, font="Aptos Display")
add_text(slide, "From first scripts to the building blocks of AI", 0.65, 0.91, 8.4, 0.25,
         12.5, (191, 211, 229))
add_text(slide, "SO FAR", 11.48, 0.49, 1.18, 0.25, 11, (191, 211, 229), True,
         align=PP_ALIGN.RIGHT)
add_text(slide, "Python basics", 10.45, 0.76, 2.2, 0.30, 16, (255, 255, 255), True,
         align=PP_ALIGN.RIGHT)

cards = [
    (0.62, 1.88, 3.85, 1.62, "01", "Core syntax", "Printed output, variables,\nstrings, numbers, and type hints.", (34, 137, 160)),
    (4.74, 1.88, 3.85, 1.62, "02", "Input + decisions", "Collected user input and used\ncomparisons with if / elif / else.", (232, 126, 67)),
    (8.86, 1.88, 3.85, 1.62, "03", "Loops + data", "Used lists with for and while\nloops to repeat useful work.", (92, 143, 94)),
    (0.62, 3.82, 3.85, 1.62, "04", "Functions + math", "Created reusable functions with\nparameters, returns, and math.sqrt().", (119, 92, 157)),
    (4.74, 3.82, 3.85, 1.62, "05", "Handling errors", "Used try / except to catch a\nTypeError and continue execution.", (191, 92, 87)),
]

for card in cards:
    add_card(slide, *card)

callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.86), Inches(3.82), Inches(3.85), Inches(1.62))
callout.fill.solid()
callout.fill.fore_color.rgb = RGBColor(229, 239, 244)
callout.line.fill.background()
add_text(slide, "READY FOR THE NEXT STEP", 9.15, 4.12, 3.20, 0.25, 11, (24, 91, 111), True)
add_text(slide, "Combine these foundations\nwith data and AI libraries.", 9.15, 4.47, 3.15, 0.55, 15, (24, 49, 77), True)

footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.78), Inches(13.333), Inches(0.72))
footer.fill.solid()
footer.fill.fore_color.rgb = RGBColor(224, 231, 238)
footer.line.fill.background()
add_text(slide, "Evidence in practice: conditional.py  |  loops.py  |  fun-sum.py  |  math_example.py  |  exception_handling.py",
         0.65, 7.00, 12.1, 0.22, 10, (65, 80, 96))

presentation.save(OUTPUT)
print(OUTPUT)